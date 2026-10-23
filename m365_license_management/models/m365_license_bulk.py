# -*- coding: utf-8 -*-
from collections import defaultdict

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from odoo import Command, _, api, fields, models
from odoo.exceptions import UserError, ValidationError

from .m365_tenant import M365GraphError

TEMPLATE_NAME = 'm365_bulk_%s_template.xlsx'  # per operation

# Columns of the file: (key, header, required, other accepted headers). The headers
# are matched case insensitively.
COLUMNS = [
    ('tenant', 'Tenant', True, ('tenant name', 'primary domain', 'domain',
                                'directory (tenant) id')),
    ('user', 'User', True, ('user principal name', 'upn', 'email', 'object id')),
    ('license', 'License', True, ('licence', 'sku', 'product', 'sku id')),
    ('reason', 'Business Justification', False, ('justification', 'reason')),
]
TITLES = {key: title for key, title, _required, _aliases in COLUMNS}
# Fields a cell is matched with (case insensitive)
TENANT_KEYS = ('name', 'domain', 'tenant_id')
USER_KEYS = ('user_principal_name', 'mail', 'graph_id')
SKU_KEYS = ('sku_part_number', 'name', 'sku_id')
# Placeholder example rows of the template per operation, used when no user and
# license of the database can be given as example. They use the example tenants of the
# tenant import template and placeholder users; each row shows another accepted way to
# name the tenant, user and license.
SAMPLE_ROWS = {
    'assign': [
        # Primary domain, user principal name, SKU
        ('contoso.com', 'adele.vance@contoso.com', 'SPE_E3', 'New joiner in Finance'),
        # Tenant name, email, product name
        ('Northwind Traders', 'megan.bowen@northwindtraders.com', 'Microsoft 365 E5',
         'Promoted to team lead'),
        # Directory (tenant) ID, object ID, SKU ID (Office 365 E3)
        ('aaaaaaaa-0000-0000-0000-000000000004', 'cccccccc-0000-0000-0000-000000000001',
         '6fd2c87f-b296-42f0-b197-1e91e994b900', 'Joined the migration project'),
    ],
    'remove': [
        ('contoso.com', 'lee.gu@contoso.com', 'ENTERPRISEPACK', 'Left the company'),
        ('Adventure Works', 'diego.siciliani@adventure-works.com',
         'Microsoft Teams Enterprise', 'License no longer needed'),
        # Microsoft 365 E3
        ('aaaaaaaa-0000-0000-0000-000000000003', 'cccccccc-0000-0000-0000-000000000002',
         '05e9a617-0261-4cee-bb44-138d3ef5d965', 'Moved to another license'),
    ],
}


class M365LicenseBulkWizard(models.TransientModel):
    """Assign or remove licenses for the users listed in an uploaded file.

    Every row (tenant, user, license) is verified first and nothing is executed
    while a row is not valid. One license request is created per row, so every
    operation follows the approval workflow and is written to the audit log.
    """
    _name = 'm365.license.bulk.wizard'
    _inherit = 'm365.file.import.mixin'
    _description = 'Bulk Assign / Remove Microsoft 365 Licenses'

    request_type = fields.Selection([
        ('assign', 'Assign License'),
        ('remove', 'Remove License'),
    ], string="Operation", required=True, default='assign')
    file = fields.Binary(string="File", attachment=False,
                         help="Excel (.xlsx) or CSV file with one tenant, user and license "
                              "per row.")
    filename = fields.Char()
    template_file = fields.Binary(attachment=False)
    state = fields.Selection([('draft', 'Draft'), ('done', 'Done')], default='draft')
    line_ids = fields.One2many('m365.license.bulk.line', 'wizard_id', string="Rows")
    line_count = fields.Integer(compute='_compute_counts')
    error_count = fields.Integer(compute='_compute_counts')
    done_count = fields.Integer(compute='_compute_counts')
    submitted_count = fields.Integer(compute='_compute_counts')
    failed_count = fields.Integer(compute='_compute_counts')

    @api.depends('line_ids.status')
    def _compute_counts(self):
        for wiz in self:
            statuses = wiz.line_ids.mapped('status')
            wiz.line_count = len(statuses)
            wiz.error_count = statuses.count('error')
            wiz.done_count = statuses.count('done')
            wiz.submitted_count = statuses.count('submitted')
            wiz.failed_count = statuses.count('failed')

    @api.onchange('file', 'request_type')
    def _onchange_file(self):
        # A new file or operation must be verified again before it is executed
        self.line_ids = [Command.clear()]

    def _reopen(self):
        return {
            'type': 'ir.actions.act_window',
            'name': _("Bulk Assign / Remove Licenses"),
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
            'context': {'dialog_size': 'extra-large'},
        }

    # ------------------------------------------------------------------
    # File reading
    # ------------------------------------------------------------------
    def _map_header(self, header):
        """Map column positions to column keys."""
        lookup = {}
        for key, title, _required, aliases in COLUMNS:
            for name in (key, title) + aliases:
                lookup[self._normalise(name)] = key
        mapping, unknown = {}, []
        for index, title in enumerate(header):
            if not title:
                continue
            key = lookup.get(self._normalise(title))
            if not key:
                unknown.append(title)
            elif key in mapping.values():
                raise UserError(_("The column \"%s\" appears more than once.", title))
            else:
                mapping[index] = key
        if unknown:
            raise UserError(_(
                "Unknown column(s): %(unknown)s.\nAllowed columns: %(allowed)s.",
                unknown=", ".join(unknown),
                allowed=", ".join(title for _key, title, _req, _aliases in COLUMNS)))
        missing = [title for key, title, required, _aliases in COLUMNS
                   if required and key not in mapping.values()]
        if missing:
            raise UserError(_("Missing required column(s): %s.", ", ".join(missing)))
        return mapping

    @staticmethod
    def _index(records, fnames, scope=None):
        """Map (id of the ``scope`` record, lower case field value) to the records."""
        index = {}
        for rec in records:
            for fname in fnames:
                if rec[fname]:
                    key = (rec[scope].id if scope else False, rec[fname].lower())
                    index[key] = index.get(key, records.browse()) | rec
        return index

    def _match(self, row, key, index, model, tenant=None):
        """Return the record named in the cell ``key`` of the row.

        An empty recordset is returned, and an error added to the row, when no or
        several records match.
        """
        text = row['texts'][key]
        if not text:
            return model
        title = TITLES[key]
        records = index.get((tenant.id if tenant else False, text.lower()), model)
        if len(records) > 1:
            row['errors'].append(_("%(column)s \"%(value)s\" matches several records: "
                                   "%(names)s.", column=title, value=text,
                                   names=", ".join(records.mapped('display_name'))))
            return model
        if not records:
            if tenant:
                row['errors'].append(_("%(column)s \"%(value)s\" not found in tenant "
                                       "%(tenant)s.", column=title, value=text,
                                       tenant=tenant.name))
            else:
                row['errors'].append(_("%(column)s \"%(value)s\" not found.",
                                       column=title, value=text))
        return records

    # ------------------------------------------------------------------
    # Verification
    # ------------------------------------------------------------------
    def _parse(self):
        """Read the uploaded file and verify every row.

        :return: list of dicts with the serial number (1 = first row below the
                 header), the cell texts, the matched tenant, user and license and
                 the error messages of each data row
        """
        self.ensure_one()
        rows = self._read_rows()
        if len(rows) < 2:
            raise UserError(_("The file contains no rows below the header."))
        mapping = self._map_header(rows[0])
        result = []
        for sn, values in enumerate(rows[1:], start=1):
            texts = dict.fromkeys(TITLES, '')
            texts.update({key: values[index] for index, key in mapping.items()
                          if index < len(values)})
            errors = [_("%s is required.", title) for key, title, required, _aliases in COLUMNS
                      if required and not texts[key]]
            result.append({'sn': sn, 'texts': texts, 'errors': errors})
        self._match_rows(result)
        self._check_rows(result)
        return result

    def _match_rows(self, rows):
        """Find the tenant, user and license of every row."""
        Tenant = self.env['m365.tenant']
        User = self.env['m365.license.user']
        Sku = self.env['m365.sku']
        tenant_index = self._index(Tenant.search([]), TENANT_KEYS)
        for row in rows:
            row['tenant'] = self._match(row, 'tenant', tenant_index, Tenant)
        tenants = Tenant.union(*(row['tenant'] for row in rows))
        domain = [('tenant_id', 'in', tenants.ids)]
        user_index = self._index(User.search_fetch(domain, USER_KEYS + ('tenant_id',)),
                                 USER_KEYS, 'tenant_id')
        sku_index = self._index(Sku.search(domain), SKU_KEYS, 'tenant_id')
        for row in rows:
            tenant = row['tenant']
            row['user'] = tenant and self._match(row, 'user', user_index, User, tenant) or User
            row['sku'] = tenant and self._match(row, 'license', sku_index, Sku, tenant) or Sku

    def _check_rows(self, rows):
        """Check that the operation can be done for every matched row."""
        assign = self.request_type == 'assign'
        users = self.env['m365.license.user'].union(*(row['user'] for row in rows))
        pending = {(request.user_id.id, request.sku_id.id)
                   for request in self.env['m365.license.request'].sudo().search([
                       ('user_id', 'in', users.ids), ('request_type', '=', self.request_type),
                       ('state', 'in', ('submitted', 'approved'))])}
        seen = {}
        requested = defaultdict(int)  # licenses to assign per SKU
        for row in rows:
            tenant, user, sku, errors = row['tenant'], row['user'], row['sku'], row['errors']
            if tenant and not tenant.allow_license_operations:
                errors.append(_("License operations are disabled for tenant %s.", tenant.name))
            elif tenant.state == 'draft':
                errors.append(_("Test the connection of tenant %s first.", tenant.name))
            if not (user and sku):
                continue
            key = (user.id, sku.id)
            if key in seen:
                errors.append(_("Duplicate row, same as SN. %s.", seen[key]))
                continue
            seen[key] = row['sn']
            has_license = sku in user.sku_ids
            if assign and has_license:
                errors.append(_("%(user)s already has %(sku)s.", user=user.name, sku=sku.name))
            elif not assign and not has_license:
                errors.append(_("%(user)s does not have %(sku)s.", user=user.name, sku=sku.name))
            elif key in pending:
                errors.append(_("A similar request is already pending for this user and "
                                "license."))
            elif assign:
                if user.last_sync and not user.usage_location:
                    errors.append(_("%s has no usage location. Microsoft requires it before a "
                                    "license can be assigned.", user.name))
                requested[sku] += 1
                if requested[sku] > sku.available_units:
                    errors.append(_("Not enough %(sku)s licenses in tenant %(tenant)s: "
                                    "%(count)s available.", sku=sku.name, tenant=tenant.name,
                                    count=sku.available_units))

    @staticmethod
    def _line_vals(row):
        return {
            'sn': row['sn'],
            'tenant_ref': row['texts']['tenant'],
            'user_ref': row['texts']['user'],
            'license_ref': row['texts']['license'],
            'reason': row['texts']['reason'] or False,
            'tenant_id': row['tenant'].id,
            'user_id': row['user'].id,
            'sku_id': row['sku'].id,
            'status': 'error' if row['errors'] else 'valid',
            'message': "\n".join(row['errors']) or False,
        }

    def action_verify(self):
        self.ensure_one()
        self.line_ids = [Command.clear()] + [Command.create(self._line_vals(row))
                                             for row in self._parse()]
        return self._reopen()

    # ------------------------------------------------------------------
    # Execution
    # ------------------------------------------------------------------
    def action_execute(self):
        self.ensure_one()
        if self.state == 'done':
            return self._reopen()
        if not self.line_ids:
            raise UserError(_("Verify the file before assigning or removing licenses."))
        # Verify again: the file, the users or the inventory may have changed since
        rows = self._parse()
        invalid = [row for row in rows if row['errors']]
        if invalid:
            raise ValidationError(_(
                "No license has been assigned or removed. Correct these rows and verify "
                "the file again:\n%s",
                "\n".join("SN. %s: %s" % (row['sn'], " ".join(row['errors']))
                          for row in invalid)))
        # Inventory is synchronised once per tenant at the end instead of after every request
        Request = self.env['m365.license.request'].with_context(m365_defer_inventory_sync=True)
        lines = []
        for row in rows:
            vals = self._line_vals(row)
            vals.update(self._submit_request(Request, row))
            lines.append(Command.create(vals))
        self.write({'state': 'done', 'line_ids': [Command.clear()] + lines})
        for tenant in self.line_ids.filtered(lambda line: line.status == 'done').tenant_id:
            try:
                tenant._sync_inventory()
            except M365GraphError:
                pass  # the operations succeeded; the scheduled sync will catch up
        return self._reopen()

    def _submit_request(self, Request, row):
        """Create and submit the request of a row; a validation error only fails its line."""
        try:
            with self.env.cr.savepoint():
                request = Request.create({
                    'tenant_id': row['tenant'].id,
                    'user_id': row['user'].id,
                    'sku_id': row['sku'].id,
                    'request_type': self.request_type,
                    'reason': row['texts']['reason'] or False,
                })
                request.action_submit()
        except UserError as exc:
            return {'status': 'failed', 'message': str(exc)}
        return {
            'request_id': request.id,
            'status': request.state if request.state in ('done', 'submitted') else 'failed',
            'message': request.error_message or False,
        }

    # ------------------------------------------------------------------
    # Template
    # ------------------------------------------------------------------
    def _template_hints(self):
        return {
            'tenant': _("Required. Tenant name, primary domain or directory (tenant) ID, "
                        "e.g. contoso.com"),
            'user': _("Required. User principal name, email or object ID, "
                      "e.g. adele.vance@contoso.com"),
            'license': _("Required. SKU, product name or SKU ID, e.g. SPE_E3. See the "
                         "Available Licenses sheet."),
            'reason': _("Optional. Stored on the license request."),
        }

    def _sample_rows(self):
        """Example rows of the template: up to 3 users and licenses of this database
        the operation can be done for, verified like the rows of an uploaded file. The
        placeholder rows are used when there is none.
        """
        assign = self.request_type == 'assign'
        samples = SAMPLE_ROWS[self.request_type]
        tenants = self.env['m365.tenant'].search([('allow_license_operations', '=', True),
                                                  ('state', '!=', 'draft')])
        skus = self.env['m365.sku'].search([('tenant_id', 'in', tenants.ids)])
        domain = [('tenant_id', 'in', tenants.ids),
                  ('account_enabled', '=', True) if assign else ('sku_ids', '!=', False)]
        rows = []
        for user in self.env['m365.license.user'].search(domain, limit=50):
            tenant = user.tenant_id
            if assign:
                options = skus.filtered(lambda s: s.tenant_id == tenant
                                        and s.available_units > 0) - user.sku_ids
            else:
                options = user.sku_ids & skus
            if not options:
                continue
            index = len(rows)
            sku = options[index % len(options)]
            # Each row names the tenant, user and license in another accepted way
            texts = [
                (tenant.domain, user.user_principal_name, sku.sku_part_number),
                (tenant.name, user.mail or user.user_principal_name, sku.name),
                (tenant.tenant_id, user.graph_id, sku.sku_id),
            ][index % 3] + (samples[index % 3][3],)
            rows.append({'sn': index + 1, 'texts': dict(zip(TITLES, texts)), 'errors': []})
        self._match_rows(rows)
        self._check_rows(rows)
        valid = [tuple(row['texts'].values()) for row in rows
                 if row['user'] and row['sku'] and not row['errors']]
        return valid[:3] or samples

    def action_download_template(self):
        self.ensure_one()
        hints = self._template_hints()
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = _("Licenses")
        required_fill = PatternFill('solid', fgColor='FFF2CC')
        for col, (key, title, required, _aliases) in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=1, column=col, value=title)
            cell.font = Font(bold=True)
            if required:
                cell.fill = required_fill
            cell.comment = Comment(hints[key], 'Odoo')
            sheet.column_dimensions[cell.column_letter].width = 36
        for sample in self._sample_rows():
            sheet.append(sample)
        sheet.freeze_panes = 'A2'
        # Reference of the licenses that can be used in the file (not read on import)
        reference = book.create_sheet(_("Available Licenses"))
        reference.append([_("Tenant"), _("Primary Domain"), _("License"), _("SKU"),
                          _("Total"), _("Assigned"), _("Available")])
        for cell in reference[1]:
            cell.font = Font(bold=True)
        tenants = self.env['m365.tenant'].search([('allow_license_operations', '=', True)])
        for sku in self.env['m365.sku'].search([('tenant_id', 'in', tenants.ids)]):
            reference.append([sku.tenant_id.name, sku.tenant_id.domain, sku.name,
                              sku.sku_part_number, sku.enabled_units, sku.consumed_units,
                              sku.available_units])
        for letter, width in zip('ABCDEFG', (24, 24, 40, 28, 10, 10, 10)):
            reference.column_dimensions[letter].width = width
        reference.freeze_panes = 'A2'
        return self._download_workbook(book, TEMPLATE_NAME % self.request_type)


class M365LicenseBulkLine(models.TransientModel):
    _name = 'm365.license.bulk.line'
    _description = 'Bulk License Operation Row'
    _order = 'sn, id'

    wizard_id = fields.Many2one('m365.license.bulk.wizard', required=True, ondelete='cascade')
    sn = fields.Integer(string="SN.")
    tenant_ref = fields.Char(string="Tenant")
    user_ref = fields.Char(string="User")
    license_ref = fields.Char(string="License")
    reason = fields.Text(string="Business Justification")
    tenant_id = fields.Many2one('m365.tenant', string="Matched Tenant")
    user_id = fields.Many2one('m365.license.user', string="Matched User")
    sku_id = fields.Many2one('m365.sku', string="Matched License")
    request_id = fields.Many2one('m365.license.request', string="Request")
    status = fields.Selection([
        ('valid', 'Valid'),
        ('error', 'Error'),
        ('done', 'Completed'),
        ('submitted', 'Waiting Approval'),
        ('failed', 'Failed'),
    ], string="Status")
    message = fields.Text(string="Message")
