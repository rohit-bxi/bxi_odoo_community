# -*- coding: utf-8 -*-
import base64
import csv
import io

import openpyxl

from odoo import _, models
from odoo.exceptions import UserError


class M365FileImportMixin(models.AbstractModel):
    """Read the rows of an uploaded CSV / Excel file and download templates.

    The inheriting wizards define the ``file``, ``filename`` and ``template_file``
    binary fields.
    """
    _name = 'm365.file.import.mixin'
    _description = 'Microsoft 365 File Import Helpers'

    @staticmethod
    def _normalise(title):
        return ' '.join(str(title).split()).lower()

    @staticmethod
    def _cell_text(value):
        if value is None:
            return ''
        if isinstance(value, bool):
            return 'yes' if value else 'no'
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        return str(value).strip()

    @staticmethod
    def _read_csv(raw):
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('latin-1')
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        return list(csv.reader(io.StringIO(text), dialect))

    @staticmethod
    def _read_xlsx(raw):
        try:
            book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - any unreadable / non-xlsx content
            raise UserError(_("The Excel file could not be read: %s", exc))
        try:
            return [list(row) for row in book.worksheets[0].iter_rows(values_only=True)]
        finally:
            book.close()

    def _read_rows(self):
        """Return the non-empty rows of the uploaded file as lists of cell texts."""
        self.ensure_one()
        if not self.file:
            raise UserError(_("Please upload a CSV or Excel (.xlsx) file."))
        raw = base64.b64decode(self.file)
        name = (self.filename or '').lower()
        if name.endswith('.xlsx'):
            rows = self._read_xlsx(raw)
        elif name.endswith('.csv'):
            rows = self._read_csv(raw)
        else:
            raise UserError(_("Unsupported file type. Please upload a .csv or .xlsx file."))
        result = []
        for row in rows:
            values = [self._cell_text(v) for v in row]
            if any(values):
                result.append(values)
        return result

    def _download_workbook(self, book, filename):
        """Store the workbook in ``template_file`` and return the download action."""
        self.ensure_one()
        out = io.BytesIO()
        book.save(out)
        self.template_file = base64.b64encode(out.getvalue())
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s/%s/template_file/%s?download=true' % (
                self._name, self.id, filename),
            'target': 'download',
        }
