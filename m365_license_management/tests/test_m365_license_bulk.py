# -*- coding: utf-8 -*-
import base64
import io
from unittest.mock import patch

import openpyxl

from odoo import fields
from odoo.exceptions import UserError, ValidationError
from odoo.tests import TransactionCase, new_test_user, tagged

from ..models.m365_license_bulk import SAMPLE_ROWS
from ..models.m365_tenant import M365GraphError
from .common import map_tenants

GUID_A = '11111111-2222-3333-4444-555555555555'
TENANT_PATH = 'odoo.addons.m365_license_management.models.m365_tenant.M365Tenant._graph_request'
HEADER = ('Tenant', 'User', 'License', 'Business Justification')


@tagged('post_install', '-at_install')
class TestM365LicenseBulk(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.tenant = cls.env['m365.tenant'].create({
            'name': 'Bulk', 'domain': 'bulk.com', 'tenant_id': 'tid-bulk',
            'client_id': 'cid', 'client_secret': 'secret', 'state': 'connected',
            'allow_license_operations': True, 'auto_approve_admins': True,
        })
        cls.sku = cls.env['m365.sku'].create({
            'tenant_id': cls.tenant.id, 'sku_id': GUID_A, 'sku_part_number': 'E3',
            'enabled_units': 2, 'consumed_units': 0,
        })
        cls.users = cls.env['m365.license.user'].create([{
            'name': 'User %s' % i, 'tenant_id': cls.tenant.id, 'graph_id': 'graph-%s' % i,
            'user_principal_name': 'user%s@bulk.com' % i, 'mail': 'user%s.mail@bulk.com' % i,
        } for i in range(4)])
        # Same SKU in another tenant: licenses are looked up in the tenant of the row
        cls.other_tenant = cls.env['m365.tenant'].create({
            'name': 'Other', 'domain': 'other.com', 'tenant_id': 'tid-other',
            'client_id': 'cid', 'client_secret': 'secret', 'state': 'connected',
            'allow_license_operations': True, 'auto_approve_admins': True,
        })
        cls.other_sku = cls.env['m365.sku'].create({
            'tenant_id': cls.other_tenant.id, 'sku_id': GUID_A, 'sku_part_number': 'E3',
            'enabled_units': 5, 'consumed_units': 0,
        })
        cls.other_user = cls.env['m365.license.user'].create({
            'name': 'Guest', 'tenant_id': cls.other_tenant.id, 'graph_id': 'graph-other',
            'user_principal_name': 'guest@other.com',
        })

    @staticmethod
    def _file(rows, header=HEADER):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.append(header)
        for row in rows:
            sheet.append(row)
        out = io.BytesIO()
        book.save(out)
        return base64.b64encode(out.getvalue())

    def _wizard(self, request_type, rows, header=HEADER, env=None):
        return (env or self.env)['m365.license.bulk.wizard'].create({
            'request_type': request_type,
            'file': self._file(rows, header),
            'filename': 'licenses.xlsx',
        })

    def _graph(self, method, path, params=None, json_body=None, paginate=False):
        if path == '/subscribedSkus':
            return [{'skuId': GUID_A, 'skuPartNumber': 'E3', 'capabilityStatus': 'Enabled',
                     'consumedUnits': 2, 'prepaidUnits': {'enabled': 2}}]
        return {}

    def _requests(self):
        return self.env['m365.license.request'].search(
            [('tenant_id', 'in', (self.tenant | self.other_tenant).ids)])

    def test_bulk_assign(self):
        first, second = self.users[1:3]
        wizard = self._wizard('assign', [
            # Tenant, user and license by domain, user principal name and SKU (any case)
            ('BULK.com', 'User1@Bulk.com', 'e3', 'New joiner'),
            # by tenant name, email and SKU ID
            ('Bulk', 'user2.mail@bulk.com', GUID_A, None),
            # by directory ID and object ID
            ('tid-other', 'graph-other', 'E3', None),
        ])
        wizard.action_verify()
        self.assertEqual((wizard.line_count, wizard.error_count), (3, 0))
        self.assertEqual(wizard.line_ids.mapped('status'), ['valid'] * 3)
        self.assertEqual(wizard.line_ids.user_id, first | second | self.other_user)
        self.assertEqual([line.sku_id for line in wizard.line_ids],
                         [self.sku, self.sku, self.other_sku])
        self.assertFalse(self._requests())
        with patch(TENANT_PATH, side_effect=self._graph) as mocked:
            wizard.action_execute()
        self.assertEqual(wizard.state, 'done')
        self.assertEqual(wizard.line_ids.mapped('status'), ['done'] * 3)
        self.assertEqual(wizard.line_ids.request_id.mapped('state'), ['done'] * 3)
        self.assertEqual(wizard.line_ids[0].request_id.reason, 'New joiner')
        self.assertEqual((first | second).sku_ids, self.sku)
        self.assertEqual(self.other_user.sku_ids, self.other_sku)
        # The inventory is synchronised once per tenant at the end
        paths = [call.args[1] for call in mocked.call_args_list]
        self.assertEqual(paths, ['/users/graph-1/assignLicense', '/users/graph-2/assignLicense',
                                 '/users/graph-other/assignLicense',
                                 '/subscribedSkus', '/subscribedSkus'])

    def test_invalid_rows_assign_nothing(self):
        self.users[0].sku_ids = self.sku
        self.users[2].last_sync = fields.Datetime.now()
        wizard = self._wizard('assign', [
            ('bulk.com', 'user0@bulk.com', 'E3'),
            ('unknown.com', 'user1@bulk.com', 'E3'),
            ('bulk.com', 'nobody@bulk.com', 'E5'),
            ('bulk.com', 'guest@other.com', 'E3'),
            ('bulk.com', None, 'E3'),
            ('bulk.com', 'user1@bulk.com', 'E3'),
            ('bulk.com', 'graph-1', GUID_A),
            ('bulk.com', 'user2@bulk.com', 'E3'),
            ('bulk.com', 'user3@bulk.com', 'E3'),
        ])
        with self.assertRaisesRegex(UserError, 'Verify the file'):
            wizard.action_execute()
        wizard.action_verify()
        self.assertEqual((wizard.line_count, wizard.error_count), (9, 8))
        self.assertEqual(wizard.line_ids.mapped('message'), [
            'User 0 already has E3.',
            'Tenant "unknown.com" not found.',
            'User "nobody@bulk.com" not found in tenant Bulk.\n'
            'License "E5" not found in tenant Bulk.',
            'User "guest@other.com" not found in tenant Bulk.',
            'User is required.',
            False,
            'Duplicate row, same as SN. 6.',
            'User 2 has no usage location. Microsoft requires it before a license can be '
            'assigned.',
            # Only 2 licenses are available
            'Not enough E3 licenses in tenant Bulk: 2 available.',
        ])
        with patch(TENANT_PATH, side_effect=self._graph) as mocked, \
                self.assertRaisesRegex(ValidationError, 'SN. 2: Tenant "unknown.com" not found'):
            wizard.action_execute()
        self.assertFalse(mocked.called)
        self.assertFalse(self._requests())

    def test_execute_verifies_the_file_again(self):
        wizard = self._wizard('assign', [('bulk.com', 'user1@bulk.com', 'E3')])
        wizard.action_verify()
        self.assertEqual(wizard.line_ids.status, 'valid')
        # Assigned by someone else in the meantime
        self.users[1].sku_ids = self.sku
        with patch(TENANT_PATH, side_effect=self._graph) as mocked, \
                self.assertRaisesRegex(ValidationError, 'SN. 1: User 1 already has E3.'):
            wizard.action_execute()
        self.assertFalse(mocked.called)
        self.assertFalse(self._requests())

    def test_tenant_not_ready(self):
        self.tenant.state = 'draft'
        self.other_tenant.allow_license_operations = False
        wizard = self._wizard('assign', [('bulk.com', 'user1@bulk.com', 'E3'),
                                         ('other.com', 'guest@other.com', 'E3')])
        wizard.action_verify()
        self.assertEqual(wizard.line_ids.mapped('message'), [
            'Test the connection of tenant Bulk first.',
            'License operations are disabled for tenant Other.',
        ])

    def test_bulk_remove_graph_failure(self):
        first, second = self.users[:2]
        (first | second).sku_ids = self.sku
        # Other accepted column titles
        header = ('Primary Domain', 'UPN', 'SKU')
        wizard = self._wizard('remove', [('bulk.com', 'user2@bulk.com', 'E3')], header)
        wizard.action_verify()
        self.assertEqual(wizard.line_ids.message, 'User 2 does not have E3.')

        def graph(method, path, params=None, json_body=None, paginate=False):
            if path == '/users/graph-1/assignLicense':
                raise M365GraphError('Request_ResourceNotFound')
            return self._graph(method, path)

        wizard = self._wizard('remove', [('bulk.com', 'user0@bulk.com', 'E3'),
                                         ('bulk.com', 'user1@bulk.com', 'E3')], header)
        wizard.action_verify()
        with patch(TENANT_PATH, side_effect=graph):
            wizard.action_execute()
        self.assertEqual(wizard.line_ids.mapped('status'), ['done', 'failed'])
        self.assertEqual(wizard.line_ids.request_id.mapped('state'), ['done', 'failed'])
        self.assertEqual(wizard.line_ids[1].message, 'Request_ResourceNotFound')
        self.assertEqual((wizard.done_count, wizard.failed_count), (1, 1))
        self.assertFalse(first.sku_ids)
        self.assertEqual(second.sku_ids, self.sku)

    def test_bulk_by_user_waits_for_approval(self):
        user = map_tenants(new_test_user(self.env, 'm365_bulk_user',
                                         groups='m365_license_management.group_m365_user'),
                           self.tenant)
        wizard = self._wizard('assign', [('bulk.com', 'user1@bulk.com', 'E3'),
                                         ('bulk.com', 'user2@bulk.com', 'E3')],
                              env=self.env(user=user))
        wizard.action_verify()
        with patch(TENANT_PATH, side_effect=self._graph) as mocked:
            wizard.action_execute()
        self.assertEqual(wizard.submitted_count, 2)
        self.assertEqual(set(wizard.line_ids.mapped('request_id.state')), {'submitted'})
        self.assertFalse(mocked.called)

    def test_bulk_by_user_only_finds_mapped_tenants(self):
        user = map_tenants(new_test_user(self.env, 'm365_bulk_user',
                                         groups='m365_license_management.group_m365_user'),
                           self.tenant)
        wizard = self._wizard('assign', [('bulk.com', 'user1@bulk.com', 'E3'),
                                         ('other.com', 'guest@other.com', 'E3')],
                              env=self.env(user=user))
        wizard.action_verify()
        self.assertEqual(wizard.line_ids.mapped('status'), ['valid', 'error'])
        self.assertEqual(wizard.line_ids[1].message, 'Tenant "other.com" not found.')

    def test_pending_and_ambiguous_rows(self):
        self.env['m365.license.request'].create({
            'tenant_id': self.tenant.id, 'user_id': self.users[1].id, 'sku_id': self.sku.id,
            'state': 'submitted'})
        self.users[2:4].write({'mail': 'shared@bulk.com'})
        # Semicolon separated CSV without the optional column
        content = 'Tenant;User;License\nbulk.com;user1@bulk.com;E3\nbulk.com;shared@bulk.com;E3\n'
        wizard = self.env['m365.license.bulk.wizard'].create({
            'request_type': 'assign', 'file': base64.b64encode(content.encode()),
            'filename': 'licenses.csv'})
        wizard.action_verify()
        self.assertEqual(wizard.line_ids.mapped('message'), [
            'A similar request is already pending for this user and license.',
            'User "shared@bulk.com" matches several records: User 2, User 3.',
        ])
        self.assertFalse(wizard.line_ids[1].user_id)

    def test_execute_once(self):
        wizard = self._wizard('assign', [('bulk.com', 'user1@bulk.com', 'E3')])
        wizard.action_verify()

        def graph(method, path, params=None, json_body=None, paginate=False):
            if path == '/subscribedSkus':
                raise M365GraphError('HTTP 503 - ServiceUnavailable')
            return {}

        with patch(TENANT_PATH, side_effect=graph):
            wizard.action_execute()
        # Assigned even if the inventory could not be synchronised afterwards
        self.assertEqual(wizard.line_ids.status, 'done')
        self.assertEqual(self.users[1].sku_ids, self.sku)
        with patch(TENANT_PATH, side_effect=graph) as mocked:
            action = wizard.action_execute()
        self.assertFalse(mocked.called)
        self.assertEqual(action['res_id'], wizard.id)
        self.assertEqual(len(self._requests()), 1)

    def test_file_errors(self):
        for header, error in ((('Tenant', 'User'), 'Missing required column'),
                              (HEADER + ('Colour',), 'Unknown column'),
                              (('Tenant', 'UPN', 'Email', 'License'), 'more than once')):
            wizard = self._wizard('assign', [('bulk.com', 'user1@bulk.com', 'E3')], header)
            with self.assertRaisesRegex(UserError, error):
                wizard.action_verify()
        wizard = self._wizard('assign', [])
        with self.assertRaisesRegex(UserError, 'no rows below the header'):
            wizard.action_verify()

    def _template(self, request_type):
        """Download the template and upload it again as is."""
        wizard = self.env['m365.license.bulk.wizard'].create({'request_type': request_type})
        action = wizard.action_download_template()
        self.assertIn('/template_file/m365_bulk_%s_template.xlsx' % request_type, action['url'])
        book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(wizard.template_file)))
        wizard.write({'file': wizard.template_file, 'filename': 'template.xlsx'})
        wizard.action_verify()
        return wizard, [list(sheet.iter_rows(values_only=True)) for sheet in book.worksheets]

    def test_download_template(self):
        # Only the tenants of the test provide example rows
        others = self.env['m365.tenant'].search([]) - self.tenant - self.other_tenant
        others.allow_license_operations = False
        self.sku.enabled_units = 5
        (self.users[0] | self.users[1]).sku_ids = self.sku
        self.other_user.sku_ids = self.other_sku
        self.env['m365.license.user'].create({
            'name': 'User 4', 'tenant_id': self.tenant.id, 'graph_id': 'graph-4',
            'user_principal_name': 'user4@bulk.com',
        })
        # Users and licenses of the database the operation can be done for, each row
        # naming them in another way
        expected = {
            'assign': [
                ('bulk.com', 'user2@bulk.com', 'E3', 'New joiner in Finance'),
                ('Bulk', 'user3.mail@bulk.com', 'E3', 'Promoted to team lead'),
                ('tid-bulk', 'graph-4', GUID_A, 'Joined the migration project'),
            ],
            'remove': [
                ('other.com', 'guest@other.com', 'E3', 'Left the company'),
                ('Bulk', 'user0.mail@bulk.com', 'E3', 'License no longer needed'),
                ('tid-bulk', 'graph-1', GUID_A, 'Moved to another license'),
            ],
        }
        for request_type, rows in expected.items():
            wizard, (licenses, reference) = self._template(request_type)
            self.assertEqual(licenses, [HEADER] + rows)
            self.assertIn(('Bulk', 'bulk.com', 'E3', 'E3', 5, 0, 5), reference)
            # The example rows are valid as they are
            self.assertEqual((wizard.line_count, wizard.error_count), (3, 0))

    def test_download_template_placeholders(self):
        # No tenant allows license operations: placeholder rows, which do not verify
        self.env['m365.tenant'].search([]).allow_license_operations = False
        for request_type in ('assign', 'remove'):
            wizard, (licenses, reference) = self._template(request_type)
            self.assertEqual(licenses, [HEADER] + SAMPLE_ROWS[request_type])
            self.assertEqual(len(reference), 1)
            self.assertEqual((wizard.line_count, wizard.error_count), (3, 3))
