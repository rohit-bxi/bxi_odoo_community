# -*- coding: utf-8 -*-
import base64
import io

import openpyxl

from odoo.exceptions import UserError
from odoo.tests import TransactionCase, tagged

from ..models.m365_tenant_import import NEW_COLUMNS, NEW_HEADERS

GUID_1 = '11111111-2222-3333-4444-555555555555'
GUID_2 = '66666666-7777-8888-9999-000000000000'
GUID_3 = 'aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee'
GUID_4 = '12345678-1234-1234-1234-123456789012'
GUID_5 = 'abcdefab-cdef-abcd-efab-cdefabcdefab'
HEADER = ('Tenant Name,Primary Domain,Directory (Tenant) ID,Application (Client) ID,'
          'Client Secret,Enable License Operations,Low Availability Threshold (%)\n')
# Customers.csv of the provisioning script (new tenants)
NEW_HEADER = ('CompanyName,Domain,FirstName,LastName,Email,PhoneNumber,AddressLine1,'
              'AddressLine2,City,State,PostalCode,Country\n')


@tagged('post_install', '-at_install')
class TestM365TenantImport(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        # The tenants of the database (e.g. imported from the template) must not clash
        # with the rows of the test files, which have to be unique among all tenants
        for tenant in cls.env['m365.tenant'].with_context(active_test=False).search([]):
            tenant.write({'name': 'Tenant %s' % tenant.id,
                          'domain': 'tenant-%s.example' % tenant.id,
                          'tenant_id': 'tenant-%s' % tenant.id})
        cls.account = cls.env['m365.partner.account'].create({
            'name': 'BXI CSP', 'partner_tenant_id': 'partner.onmicrosoft.com',
            'client_id': 'pc-app', 'default_culture': 'en-IN',
        })

    def _wizard(self, content, filename='tenants.csv', **vals):
        if isinstance(content, str):
            content = content.encode()
        return self.env['m365.tenant.import.wizard'].create(dict(vals, **{
            'file': base64.b64encode(content), 'filename': filename,
        }))

    def _new_wizard(self, content, filename='Customers.csv'):
        return self._wizard(content, filename, tenant_type='new',
                            partner_account_id=self.account.id)

    def test_import_csv(self):
        wizard = self._wizard(HEADER
                              + 'Contoso,contoso.com,%s,%s,s1,no,15\n' % (GUID_1, GUID_2)
                              + '\n'
                              + 'Fabrikam,fabrikam.com,%s,%s,,,\n' % (GUID_2.upper(), GUID_3))
        wizard.action_validate()
        self.assertEqual((wizard.line_count, wizard.error_count), (2, 0))
        # Serial numbers start at 1 below the header, blank rows are skipped
        self.assertEqual(wizard.line_ids.mapped('sn'), [1, 2])
        wizard.action_import()
        tenants = self.env['m365.tenant'].search([('tenant_id', 'in', (GUID_1, GUID_2))])
        self.assertEqual(len(tenants), 2)
        contoso = tenants.filtered(lambda t: t.name == 'Contoso')
        self.assertFalse(contoso.allow_license_operations)
        self.assertEqual(contoso.low_availability_percent, 15)
        self.assertEqual(contoso.sudo().client_secret, 's1')
        fabrikam = tenants - contoso
        # Blank or missing cells keep the model defaults (switches enabled) and GUIDs are
        # stored in lower case
        self.assertEqual(fabrikam.tenant_id, GUID_2)
        self.assertTrue(fabrikam.allow_license_operations)
        self.assertTrue(all(tenants.mapped('auto_approve_admins')))
        self.assertEqual(fabrikam.low_availability_percent, 10)
        self.assertEqual(fabrikam.graph_url, 'https://graph.microsoft.com/v1.0')
        self.assertFalse(wizard.file)

    def test_invalid_rows_create_nothing(self):
        Tenant = self.env['m365.tenant'].with_context(active_test=False)
        Tenant.create({'name': 'Existing', 'domain': 'existing.com', 'tenant_id': GUID_3,
                       'client_id': GUID_1, 'active': False})
        count = Tenant.search_count([])
        wizard = self._wizard(HEADER
                              + 'Good,good.com,%s,%s,,,\n' % (GUID_1, GUID_2)
                              + ',bad domain,not-a-guid,%s,,maybe,150\n' % GUID_2
                              + 'Dup,dup.com,%s,%s,,,\n' % (GUID_1, GUID_2)
                              + 'Old,old.com,%s,%s,,,\n' % (GUID_3.upper(), GUID_2)
                              + 'GOOD,Good.com,%s,%s,,,\n' % (GUID_4, GUID_2)
                              + 'existing,EXISTING.COM,%s,%s,,,\n' % (GUID_5, GUID_2))
        with self.assertRaisesRegex(UserError, 'Validate the file'):
            wizard.action_import()
        wizard.action_validate()
        self.assertEqual((wizard.line_count, wizard.error_count), (6, 5))
        good, bad, dup, old, same, existing = wizard.line_ids
        self.assertEqual(good.status, 'valid')
        for text in ('Tenant Name is required', 'not a valid domain', 'not a valid ID',
                     'Yes/No', 'between 0 and 100'):
            self.assertIn(text, bad.message)
        # Duplicates within the file and with existing (even archived) tenants,
        # ignoring the case; the client ID may be shared by several tenants
        self.assertEqual(dup.message, 'Duplicate Directory (Tenant) ID, same as SN. 1.')
        self.assertEqual(old.message, 'Directory (Tenant) ID "%s" already exists in Odoo.'
                         % GUID_3)
        self.assertEqual(same.message, 'Duplicate Primary Domain, same as SN. 1.\n'
                                       'Duplicate Tenant Name, same as SN. 1.')
        self.assertEqual(existing.message, 'Primary Domain "existing.com" already exists '
                                           'in Odoo.\nTenant Name "existing" already exists '
                                           'in Odoo.')
        with self.assertRaisesRegex(UserError, '(?s)SN. 2: .*SN. 6: '):
            wizard.action_import()
        self.assertEqual(Tenant.search_count([]), count)

    def test_import_checks_again(self):
        wizard = self._wizard(HEADER + 'Contoso,contoso.com,%s,%s,,,\n' % (GUID_1, GUID_2))
        wizard.action_validate()
        self.assertEqual(wizard.error_count, 0)
        # Created by someone else after the validation
        self.env['m365.tenant'].create({'name': 'Contoso Group', 'domain': 'contoso.com',
                                        'tenant_id': GUID_3, 'client_id': GUID_2})
        with self.assertRaisesRegex(UserError, 'SN. 1: Primary Domain "contoso.com"'):
            wizard.action_import()
        self.assertFalse(self.env['m365.tenant'].search([('tenant_id', '=', GUID_1)]))

    def test_invalid_values(self):
        wizard = self._wizard(
            'Tenant Name;Primary Domain;Directory (Tenant) ID;Application (Client) ID;'
            'Authority URL;Low Availability Threshold (Units);Company\n'
            'Contoso;contoso.com;%s;%s;http://login.microsoftonline.com;2.5;Unknown Co\n'
            'Fabrikam;fabrikam.com;%s;%s;;abc;\n' % (GUID_1, GUID_2, GUID_3, GUID_2)
            + 'Woodgrove;woodgrove.com;%s;%s;;-1;\n' % (GUID_4, GUID_2))
        wizard.action_validate()
        self.assertEqual(wizard.line_ids.mapped('message'), [
            'Authority URL: "http://login.microsoftonline.com" must start with https://.\n'
            'Low Availability Threshold (Units): 2.5 must be a whole number of 0 or more.\n'
            'Company: Company "Unknown Co" not found in your allowed companies.',
            'Low Availability Threshold (Units): "abc" is not a number.',
            'Low Availability Threshold (Units): -1 must be a whole number of 0 or more.',
        ])
        self.assertEqual(wizard.line_ids.mapped('directory_id'), [GUID_1, GUID_3, GUID_4])

    def test_bad_header(self):
        wizard = self._wizard('Tenant Name,Primary Domain,Tenant\nA,a.com,x\n')
        with self.assertRaisesRegex(UserError, 'Unknown column'):
            wizard.action_validate()
        wizard = self._wizard('name,domain,tenant_id\nA,a.com,%s\n' % GUID_1)
        with self.assertRaisesRegex(UserError, 'Application \\(Client\\) ID'):
            wizard.action_validate()
        with self.assertRaisesRegex(UserError, 'Unsupported file type'):
            self._wizard('x', filename='tenants.txt').action_validate()
        # The label and the technical name are the same column
        wizard = self._wizard(HEADER.replace('Client Secret', 'tenant_id') + 'A,a.com\n')
        with self.assertRaisesRegex(UserError, 'The column "tenant_id" appears more than once'):
            wizard.action_validate()
        with self.assertRaisesRegex(UserError, 'no tenant rows below the header'):
            self._wizard(HEADER).action_validate()

    def test_import_xlsx_from_template(self):
        wizard = self.env['m365.tenant.import.wizard'].create({})
        wizard.action_download_template()
        book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(wizard.template_file)))
        sheet = book.active
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(headers, [
            'Tenant Name', 'Primary Domain', 'Directory (Tenant) ID', 'Application (Client) ID',
            'Client Secret', 'Authority URL', 'Graph API URL', 'Graph Scope'])
        self.assertEqual(sheet.max_row, 4)  # header + 3 example tenants
        # Columns that are not in the template can still be added
        extra = ['Scheduled Synchronisation', 'Low Availability Threshold (Units)', 'Company']
        for col, title in enumerate(extra, start=len(headers) + 1):
            sheet.cell(row=1, column=col, value=title)
        row = dict.fromkeys(headers + extra)
        row.update({'Tenant Name': 'Woodgrove Bank', 'Primary Domain': 'woodgrovebank.com',
                    'Directory (Tenant) ID': GUID_1, 'Application (Client) ID': GUID_2,
                    'Scheduled Synchronisation': False,
                    'Low Availability Threshold (Units)': 5.0,
                    'Company': self.env.company.name})
        sheet.append(list(row.values()))
        out = io.BytesIO()
        book.save(out)
        wizard.write({'file': base64.b64encode(out.getvalue()), 'filename': 'tenants.xlsx'})
        wizard.action_validate()
        # The example tenants are valid test data
        self.assertEqual((wizard.line_count, wizard.error_count), (4, 0))
        wizard.action_import()
        Tenant = self.env['m365.tenant']
        northwind = Tenant.search([('domain', '=', 'northwindtraders.com')])
        self.assertEqual(northwind.graph_url, 'https://graph.microsoft.us/v1.0')
        adventure = Tenant.search([('domain', '=', 'adventure-works.com')])
        self.assertFalse(adventure.sudo().client_secret)
        self.assertTrue(adventure.allow_license_operations and adventure.auto_approve_admins
                        and adventure.sync_enabled and adventure.sync_users)
        tenant = Tenant.search([('tenant_id', '=', GUID_1)])
        self.assertEqual(tenant.name, 'Woodgrove Bank')
        self.assertFalse(tenant.sync_enabled)
        self.assertTrue(tenant.sync_users and tenant.allow_license_operations)
        self.assertEqual(tenant.authority_url, 'https://login.microsoftonline.com')
        self.assertEqual(tenant.low_availability_units, 5)
        self.assertEqual(tenant.company_id, self.env.company)

    # ------------------------------------------------------------------
    # New tenants (Customers.csv)
    # ------------------------------------------------------------------
    def test_import_new_tenants(self):
        wizard = self._new_wizard(
            NEW_HEADER
            + 'Fabrikam Inc,FabrikamInc,Lee,Gu,lee.gu@fabrikam.com,,1 Main St,,Redmond,WA,'
              '98052,US\n'
            + 'Northwind,northwind.onmicrosoft.com,Megan,Bowen,megan@northwind.com,,'
              '2 Pine St,Suite 5,Seattle,WA,98101,United States\n')
        wizard.action_validate()
        self.assertEqual((wizard.line_count, wizard.error_count), (2, 0))
        self.assertEqual(wizard.line_ids.mapped('domain'),
                         ['FabrikamInc', 'northwind.onmicrosoft.com'])
        action = wizard.action_import()
        self.assertIn('Check, then Create in Microsoft 365', action['params']['message'])
        tenants = self.env['m365.tenant'].search(action['params']['next']['domain'],
                                                 order='id')
        self.assertEqual(tenants.mapped('name'), ['Fabrikam Inc', 'Northwind'])
        self.assertEqual(tenants.mapped('domain'),
                         ['fabrikaminc.onmicrosoft.com', 'northwind.onmicrosoft.com'])
        self.assertEqual(set(tenants.mapped('tenant_type')), {'new'})
        self.assertEqual(set(tenants.mapped('provision_state')), {'draft'})
        self.assertEqual(set(tenants.mapped('culture')), {'en-IN'})
        self.assertEqual(tenants.country_id, self.env.ref('base.us'))
        self.assertEqual(tenants[1].address_line2, 'Suite 5')
        self.assertEqual(tenants.partner_account_id, self.account)
        self.assertFalse(any(tenants.mapped('tenant_id')))

    def test_import_new_errors(self):
        self.env['m365.tenant'].create({'tenant_type': 'new', 'name': 'Adatum',
                                        'domain': 'adatum', 'partner_account_id': self.account.id})
        wizard = self._new_wizard(
            NEW_HEADER
            + 'Adatum Again,adatum,A,B,a@adatum.com,,1 St,,Pune,MH,411001,US\n'
            + ',bad domain!,A,B,not-an-email,,1 St,,City,,,Narnia\n'
            + 'India Co,indiaco,A,B,a@india.co,,1 St,,Pune,MH,411001,IN\n'
            + 'Twice,twice,A,B,a@twice.co,,1 St,,Austin,TX,73301,US\n'
            + 'Twice Again,TWICE,A,B,a@twice.co,,1 St,,Austin,TX,73301,US\n')
        wizard.action_validate()
        self.assertEqual((wizard.line_count, wizard.error_count), (5, 4))
        messages = wizard.line_ids.mapped('message')
        self.assertEqual(messages[0], 'Primary Domain "adatum.onmicrosoft.com" already '
                                      'exists in Odoo.')
        for text in ('Country: Country "Narnia" not found', 'Tenant Name is required.',
                     'Postal Code is required.', 'Domain "bad domain!.onmicrosoft.com" is not '
                     'valid', '"not-an-email" is not a valid email address.'):
            self.assertIn(text, messages[1])
        self.assertNotIn('Country is required.', messages[1])
        self.assertEqual(messages[1].count('Tenant Name is required.'), 1)
        self.assertEqual(messages[2], 'Microsoft requires the organization registration '
                                      'number for India.')
        self.assertFalse(messages[3])
        self.assertEqual(messages[4], 'Duplicate Primary Domain, same as SN. 4.')
        with self.assertRaisesRegex(UserError, '(?s)SN. 1: .*SN. 5: '):
            wizard.action_import()
        # The Partner Center account gives the defaults of the new tenants
        wizard.partner_account_id = False
        with self.assertRaisesRegex(UserError, 'Select the Partner Center account'):
            wizard.action_validate()

    def test_import_new_headers(self):
        for header, error in (('CompanyName,Domain,Email\n', 'Missing required column'),
                              (NEW_HEADER.replace('State', 'Region'), 'Unknown column'),
                              (NEW_HEADER.replace('Email', 'Email,Contact Email'),
                               'appears more than once')):
            with self.assertRaisesRegex(UserError, error):
                self._new_wizard(header + 'x\n').action_validate()
        with self.assertRaisesRegex(UserError, 'no tenant rows'):
            self._new_wizard(NEW_HEADER).action_validate()
        # The field labels are accepted as well
        wizard = self._new_wizard('Tenant Name,Primary Domain,Contact First Name,'
                                  'Contact Last Name,Contact Email,Address Line 1,City,'
                                  'Postal Code,Country\n'
                                  'Litware,litware,Lee,Gu,lee@litware.com,1 St,Austin,73301,US\n')
        wizard.action_validate()
        self.assertEqual((wizard.line_count, wizard.error_count), (1, 0))

    def test_import_new_template(self):
        wizard = self.env['m365.tenant.import.wizard'].create(
            {'tenant_type': 'new', 'partner_account_id': self.account.id})
        action = wizard.action_download_template()
        self.assertIn('/template_file/m365_customers_template.xlsx', action['url'])
        sheet = openpyxl.load_workbook(io.BytesIO(base64.b64decode(wizard.template_file))).active
        self.assertEqual([cell.value for cell in sheet[1]],
                         [NEW_HEADERS[fname] for fname, _kind, _required in NEW_COLUMNS])
        # The example row is valid
        wizard.write({'file': wizard.template_file, 'filename': 'template.xlsx'})
        wizard.action_validate()
        self.assertEqual((wizard.line_count, wizard.error_count), (1, 0))
